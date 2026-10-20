"""
Pytest Fixtures and Test Configuration.
"""

import os
import io
import shutil
import tempfile
import pytest
from fastapi.testclient import TestClient

import pymupdf
import docx

from app.config import settings
from app.rag.embeddings import DeterministicMockEmbeddings
from app.rag.llm import MockChatModel
from app.rag.vectorstore import VectorStoreManager
from app.rag.engine import CogniStudyRAGEngine
from main import app


@pytest.fixture(scope="session", autouse=True)
def setup_test_environment():
    """Ensure tests run in mock/deterministic mode."""
    settings.DEFAULT_LLM_PROVIDER = "mock"
    settings.EMBEDDING_PROVIDER = "mock"
    settings.INDEX_STORAGE_DIR = tempfile.mkdtemp(prefix="cogni_test_storage_")
    yield
    if os.path.exists(settings.INDEX_STORAGE_DIR):
        shutil.rmtree(settings.INDEX_STORAGE_DIR, ignore_errors=True)


@pytest.fixture
def mock_embeddings():
    return DeterministicMockEmbeddings(dimension=128)


@pytest.fixture
def mock_llm():
    return MockChatModel()


@pytest.fixture
def temp_vectorstore(mock_embeddings):
    temp_dir = tempfile.mkdtemp(prefix="cogni_vs_test_")
    mgr = VectorStoreManager(base_storage_dir=temp_dir)
    mgr.embedding_provider = mock_embeddings
    yield mgr
    if os.path.exists(temp_dir):
        shutil.rmtree(temp_dir, ignore_errors=True)


@pytest.fixture
def rag_engine(temp_vectorstore, mock_llm):
    return CogniStudyRAGEngine(
        vectorstore_manager=temp_vectorstore,
        llm=mock_llm
    )


@pytest.fixture
def client():
    return TestClient(app)


@pytest.fixture
def sample_pdf_bytes():
    """Generates an in-memory PDF with 2 pages of test content."""
    doc = pymupdf.open()
    page1 = doc.new_page()
    page1.insert_text((50, 50), "Page 1: Retrieval-Augmented Generation (RAG) is a machine learning architecture.")
    page2 = doc.new_page()
    page2.insert_text((50, 50), "Page 2: Vector embeddings project semantic meaning into dense mathematical vector spaces.")
    pdf_bytes = doc.tobytes()
    doc.close()
    return pdf_bytes


@pytest.fixture
def sample_docx_bytes():
    """Generates an in-memory DOCX document."""
    doc = docx.Document()
    doc.add_heading("Deep Learning Notes", level=1)
    doc.add_paragraph("Neural networks learn hierarchical representations through backpropagation and gradient descent.")
    stream = io.BytesIO()
    doc.save(stream)
    return stream.getvalue()


@pytest.fixture
def sample_txt_bytes():
    content = "Distributed computing frameworks scale data processing across multiple worker clusters using map-reduce algorithms."
    return content.encode("utf-8")
